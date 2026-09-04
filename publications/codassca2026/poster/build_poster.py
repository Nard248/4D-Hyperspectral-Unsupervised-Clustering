"""Build the CODASSCA 2026 A0 poster as an editable PowerPoint file.

    .venv/bin/python publications/codassca2026/poster/build_poster.py

All positions are in POINTS on an A0 portrait sheet (2383.9 x 3370.4 pt = 841 x 1189 mm),
mirroring the layout of the previous Publisher poster (Calibri 22 pt body, Calibri Bold 40 pt
section bars, Arial Bold title, sage-green / khaki-gold bars, navy table header).
Edit the CONTENT block to change any text; re-run to rebuild.
"""
from pathlib import Path
from PIL import Image
from lxml import etree
from pptx import Presentation
from pptx.util import Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
FIG = ROOT / "figures"
OUT = ROOT / "CODASSCA2026_poster_A0.pptx"
PAGE_W, PAGE_H = 2383.94, 3370.39                    # A0 portrait in points
M, GUT = 36.0, 24.0                                  # page margin, column gutter
CW = PAGE_W - 2 * M                                  # content width (2311.9)
COL3 = (CW - 2 * GUT) / 3                            # three-column width (754.6)

BLACK, INK, INK2 = "000000", "1A2332", "52514E"
GOLD, SAGE, SAGE_LIGHT = "E2D5A3", "C4D0B5", "D8E0CE"
ROW_BLUE, ROW_PEACH, WHITE = "EAF1F9", "FBEFEA", "FFFFFF"
RED = "C0392B"

# ------------------------------------------------------------------ text content (edit here)
SUP = {"sup": True}
CONTENT = {
    "title": ["Label-Free, Dependency-Aware Channel Selection", "for Coupled Multi-Channel Sensor Data"],
    "authors": [("Narek Meloyan", {"underline": True}), ("1,2,3", SUP), (", Narine Sarvazyan", {}), ("1,2,3,4", SUP)],
    "affiliations": [
        [("1", SUP), ("Zaven P. and Sonia Akian College of Science and Engineering, American University of Armenia, Yerevan, Armenia", {})],
        [("2", SUP), ("Engineering Research Center, American University of Armenia, Yerevan, Armenia", {})],
        [("3", SUP), ("L.A. Orbeli Institute of Physiology NAS RA, Yerevan, Armenia      ", {}), ("4", SUP), ("George Washington University, Washington, DC, USA", {})],
    ],
    "emails": "narek_meloyan@edu.aua.am   ·   phynas@gwu.edu",
    "intro": [
        "Complex sensors deliver dozens to hundreds of channels that overlap heavily and arrive in natural groups: inertial units, excitation wavelengths, scalp regions, process subsystems.",
        "For cost, power, bandwidth and interpretability we often want to keep a handful of the actual channels rather than a learned mixture of them.",
        [("The hard part is deciding ", {}), ("which", {"bold": True}), (" channels to keep when no class labels are available.", {})],
    ],
    "problem": [
        [("Unsupervised filters such as variance ranking or PCA loadings score each channel by its own ", {}), ("marginal", {"bold": True}),
         (" statistics: fast, but blind to how channels depend on one another, and biased toward the loudest signal, which is not always the informative one.", {})],
        "Supervised dependency-aware selectors (mRMR, JMI) model conditional relevance but need labels.",
        [("Few methods are at once ", {}), ("label-free, dependency-aware and discrete", {"bold": True}), (", returning real channels rather than a projection.", {})],
    ],
    "novelty": [
        [("A label-free, dependency-aware selector that returns real channels. ", {"bold": True}),
         ("A group-structured autoencoder with a ", {}), ("compressive bottleneck", {"bold": True}),
         (" learns the joint structure of all channels; ", {}), ("perturbing each latent factor", {"bold": True}),
         (" and measuring the per-channel reconstruction change gives a label-free relevance; a ", {}), ("relevance–redundancy rule", {"bold": True}),
         (" turns it into an ordered shortlist. Only the encoder changes between modalities. Its channels match a supervised selector and beat PCA and variance where the discriminative signal is not the loudest.", {})],
    ],
    "method_bar": "Method: compressive group autoencoder  →  perturbation attribution  →  relevance–redundancy selection",
    "cap_fig1": [("Figure 1: ", {"bold": True}), ("The three-stage pipeline with real intermediate outputs from the controlled benchmark: per-group encoders and mean fusion give a compressive code z; shifting each latent coordinate and decoding attributes a reconstruction change to every channel; MMR orders the channels.", {})],
    "cap_fig2": [("Figure 2: ", {"bold": True}), ("Attribution on the benchmark: latent coordinates ranked by variance (1), per-channel reconstruction change of each coordinate (2), accumulated relevance with the eight MMR picks (3).", {})],
    "stages": [
        ("1 · Compressive group autoencoder",
         "Each group is encoded independently (Conv1D for signals, Conv2D for images); the features are averaged into a latent code z with no time or space extent and decoded per group. Training uses reconstruction error only.",
         "eq1.png"),
        ("2 · Perturbation attribution",
         "Latent coordinates are ranked by variance and the top m kept; each is shifted by magnitudes δ, and the accumulated per-channel change of the decoded output, weighted by its variance rank w_j, is the label-free relevance.",
         "eq2.png"),
        ("3 · Relevance–redundancy selection",
         "Relevances are normalised within each group; channels are picked greedily by maximal marginal relevance (relevance minus λ × cosine similarity to channels already picked). The ordered output truncates to any budget K.",
         "eq3.png"),
    ],
    "protocol": [("Protocol: ", {"bold": True, "italic": True}), ("selection is always unsupervised. Labels are used only by a downstream classifier that scores the selected K channels (accuracy or macro-F1) against variance ranking, PCA loadings, a supervised mutual-information (MI) selector and random-K.", {"italic": True})],
    "syn_bar": "Controlled benchmark: four latent factors hidden among 64 channels",
    "syn_caps": [
        [("(a) ", {"bold": True}), ("64 grouped time-series channels: four independent latent factors, each in a cluster of five correlated channels of decreasing amplitude, plus 44 pure-noise channels. The class (16 = 2⁴ factor states) is decodable only by covering all four factors.", {})],
        [("(b) ", {"bold": True}), ("kNN accuracy vs K. Eight label-free channels beat the full noisy set (0.94 vs 0.87) and match the supervised selector; PCA also recovers the factors, which here are the loudest channels; variance piles into one cluster (0.24).", {})],
        [("(c) ", {"bold": True}), ("The channels each selector keeps at K = 8: the label-free method covers all four factors with no noise channel. (The bottleneck occasionally under-encodes a factor: 0.70 ± 0.29 mean accuracy over random draws of the benchmark.)", {})],
    ],
    "eeg_bar": "Motor-imagery EEG (BCI Competition IV-2a): the informative rhythm is not the loudest signal",
    "eeg_facts": [
        [("22 electrodes", {"bold": True}), (" in 4 scalp-region groups", {})],
        [("4 classes: ", {"bold": True}), ("left hand, right hand, feet, tongue", {})],
        [("9 subjects", {"bold": True}), (" · 250 Hz · band-pass 8–30 Hz", {})],
        [("Protocol: ", {"bold": True}), ("selection + CSP/LDA fit on session 1, test on session 2 (within-subject)", {})],
    ],
    "eeg_caps": [
        [("(a) ", {"bold": True}), ("Montage and groups: the grouping is acquisition metadata, not labels.", {})],
        [("(b) ", {"bold": True}), ("Why marginal methods fail here: signal power peaks over lateral and occipital electrodes (muscle, alpha, ocular activity), so variance and PCA pick loud but uninformative electrodes (C5, C6, POz); the reconstruction-based relevance concentrates on the sensorimotor strip (FC3, Cz, C1, CP1, CP3).", {})],
        [("(c) ", {"bold": True}), ("Macro-F1 vs K, mean over nine subjects. At six of 22 electrodes the label-free method reaches 0.517 (89% of the full 0.578), above PCA, variance and random and close to supervised MI.", {})],
    ],
    "hsi_bar": "Biomedical multi-excitation fluorescence imaging: up to 95% of the bands removed without loss",
    "hsi_caps": [
        [("(a) ", {"bold": True}), ("Multi-excitation fluorescence cubes: for each excitation wavelength a full emission spectrum is recorded per pixel, so every excitation–emission pair is one image and one channel: 192 bands for lichens (4 classes), 158 for collagen sponges (3 classes). The band images shown are among the selected ones.", {})],
        [("(b) ", {"bold": True}), ("Where the selected bands sit on the excitation × emission grid (background: mean emission per excitation, row-normalised). The picks spread across excitations instead of clustering on the brightest bands.", {})],
        [("(c) ", {"bold": True}), ("9 of 192 lichen bands: 89.4% vs 88.2% with all bands (a 95% reduction); 30 of 158 collagen bands: 85.6% vs 79.8%. Pruning also denoises.", {})],
    ],
    "res_bar": "Results across regimes: where label-free selection pays, and where nothing beats a random subset",
    "table": [
        ["Selector", "Synthetic\naccuracy, K = 8 / 64", "EEG\nmacro-F1, K = 6 / 22"],
        ["All channels (full set)", "0.87", "0.58"],
        ["Supervised MI", "0.99", "0.53"],
        [("Proposed (label-free)", {"bold": True, "color": RED}), ("0.94", {"bold": True, "color": RED}), ("0.52", {"bold": True, "color": RED})],
        ["PCA", "0.95", "0.48"],
        ["Variance", "0.24", "0.47"],
        ["Random-K", "0.25", "0.49"],
    ],
    "cap_table": [("Table 1: ", {"bold": True}), ("Downstream classifier score on the selected channels; no labels are used for selection. The label-free method stays within 0.05 (synthetic) and 0.01 (EEG) of the supervised selector.", {})],
    "cap_fig3": [("Figure 3: ", {"bold": True}), ("Gain over a random subset of the same size. Selection pays where the informative channels are a small, non-redundant fraction of the input; where channels are near-duplicates (wearable HAR, remote sensing) not even the supervised selector beats random, and a label-using check flags this in advance.", {})],
    "conclusion": [
        "•  A compressive group autoencoder, perturbation attribution and a relevance–redundancy rule select real channels without labels; only the encoder changes between modalities.",
        "•  The selected channels match a supervised selector and beat PCA and variance where the discriminative signal is not the loudest; up to 95% of imaging bands removed without loss.",
        "•  A simple label-using check tells in advance when any selection is pointless.",
        [("Next: ", {"bold": True}), ("tighten the bottleneck, make the attribution less sensitive to channel variance, scale to larger channel arrays.", {})],
    ],
    "references": [
        "[1] Balın, Abid, Zou. Concrete autoencoders: differentiable feature selection and reconstruction. ICML 2019.",
        "[2] Peng, Long, Ding. Feature selection based on mutual information (mRMR). IEEE TPAMI 27(8), 2005.",
        "[3] Brown, Pocock, Zhao, Luján. Conditional likelihood maximisation (JMI). JMLR 13, 2012.",
        "[4] He, Cai, Niyogi. Laplacian score for feature selection. NeurIPS 2005.",
        "[5] Cai, Liu, Cai. BS-Nets: end-to-end band selection of hyperspectral images. IEEE TGRS 58(3), 2020.",
        "[6] Rajabinasab, Houle, Chelly, Zimek. Worse than random? A baseline for unsupervised feature selection. arXiv:2605.22973, 2026.",
        "[7] Tangermann et al. Review of the BCI Competition IV. Front. Neurosci. 6, 2012.",
    ],
    "ack": "Financial support from the European Union NAS SAR-101087403 and the United States R44 HL120511 awards is gratefully acknowledged. Narek Chilingaryan and Kristina Ghahramanyan are thanked for providing the labelled lichen and collagen-sponge datasets.",
    "ack_links": [("Paper: ", {"bold": True}), ("DOI 10.1109/CODASSCA69992.2026.00059", {}), ("     Code: ", {"bold": True}), ("github.com/Nard248/spectral-select", {})],
}

# ------------------------------------------------------------------ helpers
prs = Presentation()
prs.slide_width, prs.slide_height = Emu(int(PAGE_W * 12700)), Emu(int(PAGE_H * 12700))
slide = prs.slides.add_slide(prs.slide_layouts[6])


def E(pt):
    return Emu(int(round(pt * 12700)))


def rgb(h):
    return RGBColor.from_string(h)


def add_rect(x, y, w, h, fill):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, E(x), E(y), E(w), E(h))
    s.fill.solid(); s.fill.fore_color.rgb = rgb(fill); s.line.fill.background(); s.shadow.inherit = False
    return s


def add_text(x, y, w, h, paras, size=22, font="Calibri", color=BLACK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
             spacing=1.15, space_after=5, margin=4, name=None):
    """paras: list of paragraphs; a paragraph is a string or a list of runs; a run is a string or (text, opts)."""
    tb = slide.shapes.add_textbox(E(x), E(y), E(w), E(h))
    if name: tb.name = name
    tf = tb.text_frame; tf.word_wrap = True; tf.auto_size = MSO_AUTO_SIZE.NONE; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = E(margin); tf.margin_top = tf.margin_bottom = E(margin / 2)
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.line_spacing = spacing; p.space_after = Pt(space_after)
        runs = para if isinstance(para, list) else [para]
        for r in runs:
            text, o = (r, {}) if isinstance(r, str) else r
            run = p.add_run(); run.text = text; f = run.font
            f.name = o.get("font", font); f.size = Pt(o.get("size", size))
            f.bold = o.get("bold", False); f.italic = o.get("italic", False)
            f.color.rgb = rgb(o.get("color", color))
            if o.get("underline"): f.underline = True
            if o.get("sup"): f._element.set("baseline", "30000")
    return tb


def section(x, y, w, title, fill, h=60, size=40):
    add_rect(x, y, w, h, fill)
    add_text(x, y, w, h, [[(title, {"bold": True, "size": size})]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
             spacing=1.0, space_after=0)


def add_image(path, x, y, w, h, halign="center", valign="middle", scale=1.0):
    iw, ih = Image.open(path).size
    s = min(w / iw, h / ih) * scale; dw, dh = iw * s, ih * s
    ox = x + (w - dw) / 2 if halign == "center" else (x if halign == "left" else x + w - dw)
    oy = y + (h - dh) / 2 if valign == "middle" else (y if valign == "top" else y + h - dh)
    return slide.shapes.add_picture(str(path), E(ox), E(oy), E(dw), E(dh))


def add_equation(name, x, y, w, h, scale=1.55):
    """LaTeX PNG rendered at 600 dpi (12 pt \\Large ≈ 14.4 pt): scale ~1.55 gives ≈ 22-pt maths on the sheet."""
    p = FIG / name; iw, ih = Image.open(p).size
    dw, dh = iw / 600 * 72 * scale, ih / 600 * 72 * scale
    if dw > w: dh *= w / dw; dw = w
    if dh > h: dw *= h / dh; dh = h
    return slide.shapes.add_picture(str(p), E(x + (w - dw) / 2), E(y + (h - dh) / 2), E(dw), E(dh))


def set_cell_border(cell, color=WHITE, width_pt=2.0):
    tcPr = cell._tc.get_or_add_tcPr()
    for i, tag in enumerate(("a:lnL", "a:lnR", "a:lnT", "a:lnB")):
        old = tcPr.find(qn(tag))
        if old is not None: tcPr.remove(old)
        ln = etree.Element(qn(tag), w=str(int(width_pt * 12700)), cap="flat", cmpd="sng", algn="ctr")
        sf = etree.SubElement(ln, qn("a:solidFill")); etree.SubElement(sf, qn("a:srgbClr"), val=color)
        etree.SubElement(ln, qn("a:prstDash"), val="solid")
        tcPr.insert(i, ln)


def add_table(x, y, col_widths, rows, row_h=33, size=22, highlight_row=3):
    n_r, n_c = len(rows), len(rows[0])
    gf = slide.shapes.add_table(n_r, n_c, E(x), E(y), E(sum(col_widths)), E(row_h * n_r)); tbl = gf.table
    tblPr = tbl._tbl.tblPr; tblPr.set("firstRow", "0"); tblPr.set("bandRow", "0")
    sid = tblPr.find(qn("a:tableStyleId"))
    if sid is not None: sid.text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"       # "No Style, No Grid"
    for j, cw in enumerate(col_widths): tbl.columns[j].width = E(cw)
    for i in range(n_r): tbl.rows[i].height = E(row_h if i else row_h * 1.75)
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            text, o = (val, {}) if isinstance(val, str) else val
            cell = tbl.cell(i, j); cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(INK if i == 0 else (ROW_PEACH if i == highlight_row else (ROW_BLUE if i % 2 else WHITE)))
            cell.margin_left = cell.margin_right = E(10); cell.margin_top = cell.margin_bottom = E(2)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame; tf.word_wrap = True
            for k, line in enumerate(text.split("\n")):
                p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
                p.alignment = PP_ALIGN.LEFT if j == 0 else PP_ALIGN.CENTER; p.space_after = Pt(0); p.line_spacing = 1.0
                run = p.add_run(); run.text = line; f = run.font; f.name = "Calibri"
                f.size = Pt(o.get("size", size if (i or k == 0) else size - 3)); f.bold = o.get("bold", i == 0)
                f.color.rgb = rgb(WHITE if i == 0 else o.get("color", INK))
            set_cell_border(cell)
    return gf


# ------------------------------------------------------------------ header
add_image(FIG / "logo_aua.png", 44, 118, 350, 215, halign="left")
add_image(FIG / "logo_orbeli.png", 2094, 56, 250, 250)
add_text(400, 28, 1584, 152, [[(t, {"bold": True, "size": 64, "font": "Arial"})] for t in CONTENT["title"]],
         align=PP_ALIGN.CENTER, spacing=1.0, space_after=0, name="Title")
add_text(400, 184, 1584, 52, [[(t, dict(o, font="Arial", size=38)) for t, o in CONTENT["authors"]]],
         align=PP_ALIGN.CENTER, spacing=1.0, space_after=0, name="Authors")
add_text(300, 238, 1784, 104, [[(t, dict(o, font="Arial", size=24)) for t, o in line] for line in CONTENT["affiliations"]],
         align=PP_ALIGN.CENTER, spacing=1.02, space_after=0, name="Affiliations")
add_text(400, 340, 1584, 30, [[(CONTENT["emails"], {"font": "Arial", "size": 20, "color": INK2})]], align=PP_ALIGN.CENTER,
         spacing=1.0, space_after=0, name="Emails")

# ------------------------------------------------------------------ row 1: introduction / problem / novelty
Y1 = 380
for i, (title, key, fill) in enumerate([("Introduction", "intro", SAGE_LIGHT), ("Problem Statement", "problem", GOLD), ("Novelty", "novelty", GOLD)]):
    x = M + i * (COL3 + GUT)
    section(x, Y1, COL3, title, fill)
    add_text(x, Y1 + 70, COL3, 232, CONTENT[key], size=22, align=PP_ALIGN.JUSTIFY, spacing=1.04, space_after=7, name=f"Text {title}")

# ------------------------------------------------------------------ row 2: method
Y2 = 690
section(M, Y2, CW, CONTENT["method_bar"], GOLD)
FY, FH = Y2 + 74, 336
add_image(FIG / "fig_pipeline.png", M, FY, 1334, FH, halign="left")
add_image(FIG / "fig_attribution.png", M + 1334 + GUT, FY, CW - 1334 - GUT, FH)
add_text(M, FY + FH + 4, 1334, 60, [CONTENT["cap_fig1"]], size=20, align=PP_ALIGN.JUSTIFY, spacing=1.0, name="Caption Fig1")
add_text(M + 1334 + GUT, FY + FH + 4, CW - 1334 - GUT, 60, [CONTENT["cap_fig2"]], size=20, align=PP_ALIGN.JUSTIFY, spacing=1.0, name="Caption Fig2")
SY = FY + FH + 66
for i, (head, body, eq) in enumerate(CONTENT["stages"]):
    x = M + i * (COL3 + GUT)
    add_text(x, SY, COL3, 112, [[(head, {"bold": True, "size": 23})], body], size=20, align=PP_ALIGN.JUSTIFY, spacing=1.0, space_after=3, name=f"Stage {i+1}")
    add_equation(eq, x + 10, SY + 114, COL3 - 20, 56)
add_text(M, SY + 172, CW, 30, [CONTENT["protocol"]], size=20, align=PP_ALIGN.CENTER, spacing=1.0, space_after=0, name="Protocol")

# ------------------------------------------------------------------ dataset rows
def dataset_row(y, bar_title, panels, caps, cap_h=82, panel_h=301):
    section(M, y, CW, bar_title, SAGE)
    py = y + 72
    for (img, x, w) in panels:
        add_image(FIG / img, x, py, w, panel_h)
    for (text, x, w) in caps:
        add_text(x, py + panel_h + 3, w, cap_h, [text], size=20, align=PP_ALIGN.JUSTIFY, spacing=1.0, space_after=0)
    return py + panel_h + 4 + cap_h

Y3 = 1374
p_a, p_b, p_c = (M, 780), (M + 780 + GUT, 700), (M + 780 + GUT + 700 + GUT, CW - 780 - 700 - 2 * GUT)
dataset_row(Y3, CONTENT["syn_bar"],
            [("fig_syn_dataset.png", *p_a), ("fig_syn_curve.png", *p_b), ("fig_syn_selection.png", *p_c)],
            [(CONTENT["syn_caps"][0], *p_a), (CONTENT["syn_caps"][1], *p_b), (CONTENT["syn_caps"][2], *p_c)])

Y4 = Y3 + 462
xa, wa = M, 470; xf, wf = M + 470 + 6, 238; xb, wb = xf + wf + GUT, 900; xc, wc = xb + wb + GUT, CW - (xb + wb + GUT - M)
section(M, Y4, CW, CONTENT["eeg_bar"], SAGE)
add_image(FIG / "fig_eeg_montage.png", xa, Y4 + 72, wa, 301, halign="left")
add_text(xf, Y4 + 74, wf, 301, CONTENT["eeg_facts"], size=17, spacing=1.0, space_after=7, name="EEG facts")
add_image(FIG / "fig_eeg_topomaps.png", xb, Y4 + 72, wb, 301)
add_image(FIG / "fig_eeg_curve.png", xc, Y4 + 72, wc, 301)
cy = Y4 + 72 + 301 + 3
add_text(xa, cy, wa + wf + 8, 82, [CONTENT["eeg_caps"][0]], size=20, align=PP_ALIGN.JUSTIFY, spacing=1.0, space_after=0)
add_text(xb, cy, wb, 82, [CONTENT["eeg_caps"][1]], size=20, align=PP_ALIGN.JUSTIFY, spacing=1.0, space_after=0)
add_text(xc, cy, wc, 82, [CONTENT["eeg_caps"][2]], size=20, align=PP_ALIGN.JUSTIFY, spacing=1.0, space_after=0)

Y5 = Y4 + 462
h_a, h_b, h_c = (M, 1180), (M + 1180 + GUT, 640), (M + 1180 + GUT + 640 + GUT, CW - 1180 - 640 - 2 * GUT)
dataset_row(Y5, CONTENT["hsi_bar"],
            [("fig_hsi_dataset.png", *h_a), ("fig_hsi_eem.png", *h_b), ("fig_hsi_bars.png", *h_c)],
            [(CONTENT["hsi_caps"][0], *h_a), (CONTENT["hsi_caps"][1], *h_b), (CONTENT["hsi_caps"][2], *h_c)])

# ------------------------------------------------------------------ row 6: results table + boundary chart
Y6 = Y5 + 462
section(M, Y6, CW, CONTENT["res_bar"], SAGE)
TY = Y6 + 72
add_table(M, TY, [400, 240, 240], CONTENT["table"], row_h=27, size=20)
add_text(M, TY + 212, 880, 58, [CONTENT["cap_table"]], size=20, align=PP_ALIGN.JUSTIFY, spacing=1.0, space_after=0, name="Caption Table")
chart_x = M + 880 + GUT; chart_w = CW - 880 - GUT
add_image(FIG / "fig_boundary.png", chart_x, TY, chart_w, 222, valign="top")
add_text(chart_x, TY + 222, chart_w, 52, [CONTENT["cap_fig3"]], size=18, align=PP_ALIGN.JUSTIFY, spacing=1.0, space_after=0, name="Caption Fig3")

# ------------------------------------------------------------------ footer
Y7 = Y6 + 72 + 262 + 14
FT = Y7 + 66
for i, (title, fill) in enumerate([("Conclusion and Next Steps", SAGE), ("References", SAGE), ("Acknowledgments", SAGE)]):
    section(M + i * (COL3 + GUT), Y7, COL3, title, fill, h=56)
add_text(M, FT, COL3, 169, CONTENT["conclusion"], size=17, spacing=1.0, space_after=4, align=PP_ALIGN.JUSTIFY, name="Conclusion")
add_text(M + COL3 + GUT, FT, COL3, 169, CONTENT["references"], size=15, spacing=1.0, space_after=2, name="References")
ax_ = M + 2 * (COL3 + GUT)
add_text(ax_, FT, COL3 - 210, 120, [CONTENT["ack"]], size=18, spacing=1.0, space_after=6, align=PP_ALIGN.JUSTIFY, name="Acknowledgments")
add_text(ax_, FT + 112, COL3 - 210, 56, [CONTENT["ack_links"]], size=16, spacing=1.0, space_after=0, name="Links")
add_image(FIG / "qr_paper.png", ax_ + COL3 - 200, FT + 2, 94, 94)
add_text(ax_ + COL3 - 205, FT + 96, 104, 20, [[("paper", {"size": 15, "color": INK2})]], align=PP_ALIGN.CENTER, spacing=1.0, space_after=0)
add_image(FIG / "qr_code.png", ax_ + COL3 - 98, FT + 2, 94, 94)
add_text(ax_ + COL3 - 103, FT + 96, 104, 20, [[("code", {"size": 15, "color": INK2})]], align=PP_ALIGN.CENTER, spacing=1.0, space_after=0)

prs.save(OUT)
print("wrote", OUT, "| footer bottom =", round(FT + 169), "of", round(PAGE_H))
