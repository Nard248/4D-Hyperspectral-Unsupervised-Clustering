"""Minimal Markdown -> Word (.docx) converter for the SpectraForge reports. Handles the subset of
Markdown used in docs/spectraforge/*.md: ATX headings (#..####), paragraphs, bullet/numbered lists,
blockquotes, horizontal rules, pipe tables, and inline **bold** / `code`.

Usage:  python reports/md_to_docx.py <input.md> <output.docx>
"""
from __future__ import annotations

import os
import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

_IMG = re.compile(r"^!\[(.*?)\]\((.*?)\)\s*$")

_INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`)")


def _add_runs(paragraph, text):
    for tok in _INLINE.split(text):
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**"):
            run = paragraph.add_run(tok[2:-2])
            run.bold = True
        elif tok.startswith("`") and tok.endswith("`"):
            run = paragraph.add_run(tok[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(0xB0, 0x30, 0x60)
        else:
            paragraph.add_run(tok)


def _is_table_row(line):
    return line.strip().startswith("|") and line.strip().endswith("|")


def _split_row(line):
    return [c.strip() for c in line.strip().strip("|").split("|")]


def convert(md_path, docx_path):
    with open(md_path, encoding="utf-8") as f:
        lines = f.read().splitlines()

    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)

    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        # image: ![caption](path)
        mimg = _IMG.match(stripped)
        if mimg:
            cap, rel = mimg.group(1), mimg.group(2)
            path = rel if os.path.isabs(rel) else os.path.join(os.path.dirname(os.path.abspath(md_path)), rel)
            if os.path.exists(path):
                p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(path, width=Inches(6.0))
                if cap:
                    c = doc.add_paragraph(); c.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    r = c.add_run(cap); r.italic = True; r.font.size = Pt(9)
            i += 1
            continue

        # horizontal rule
        if re.fullmatch(r"-{3,}|\*{3,}", stripped):
            doc.add_paragraph().add_run("_" * 60).font.color.rgb = RGBColor(0xCC, 0xCC, 0xCC)
            i += 1
            continue

        # headings
        m = re.match(r"(#{1,6})\s+(.*)", stripped)
        if m:
            level = len(m.group(1))
            text = m.group(2).strip()
            if level == 1:
                doc.add_heading(text, level=0)
            else:
                h = doc.add_heading(level=min(level - 1, 4))
                _add_runs(h, text)
            i += 1
            continue

        # tables
        if _is_table_row(line) and i + 1 < len(lines) and re.search(r"-{2,}", lines[i + 1]):
            header = _split_row(line)
            body = []
            i += 2  # skip header + separator
            while i < len(lines) and _is_table_row(lines[i]):
                body.append(_split_row(lines[i]))
                i += 1
            table = doc.add_table(rows=1, cols=len(header))
            table.style = "Light Grid Accent 1"
            for j, cell in enumerate(table.rows[0].cells):
                cell.paragraphs[0].add_run(header[j].replace("**", "")).bold = True
            for row in body:
                cells = table.add_row().cells
                for j in range(min(len(cells), len(row))):
                    _add_runs(cells[j].paragraphs[0], row[j])
            doc.add_paragraph()
            continue

        # blockquote
        if stripped.startswith(">"):
            p = doc.add_paragraph(style="Intense Quote")
            _add_runs(p, stripped.lstrip("> ").strip())
            i += 1
            continue

        # bullet list
        if re.match(r"[-*]\s+", stripped):
            p = doc.add_paragraph(style="List Bullet")
            _add_runs(p, re.sub(r"^[-*]\s+", "", stripped))
            i += 1
            continue

        # numbered list
        if re.match(r"\d+\.\s+", stripped):
            p = doc.add_paragraph(style="List Number")
            _add_runs(p, re.sub(r"^\d+\.\s+", "", stripped))
            i += 1
            continue

        # plain paragraph
        p = doc.add_paragraph()
        _add_runs(p, stripped)
        i += 1

    doc.save(docx_path)
    print(f"wrote {docx_path}")


if __name__ == "__main__":
    convert(sys.argv[1], sys.argv[2])
